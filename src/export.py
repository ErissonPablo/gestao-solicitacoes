# -*- coding: utf-8 -*-
"""Relatorio Excel com todas as visoes, pronto para mandar por e-mail."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AZUL = "1C5CAB"


def _formatar(ws, df: pd.DataFrame):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=AZUL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for i, col in enumerate(df.columns, start=1):
        amostra = [len(str(v)) for v in df[col].head(200).tolist()]
        largura = min(max([len(str(col))] + amostra) + 2, 55)
        ws.column_dimensions[get_column_letter(i)].width = largura
        nome = str(col).upper()
        fmt = None
        if "R$" in nome or nome.startswith("VALOR"):
            fmt = '#,##0.00'
        elif pd.api.types.is_datetime64_any_dtype(df[col]):
            fmt = "DD/MM/YYYY"
        if fmt:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = fmt


def relatorio_excel(abas: dict[str, pd.DataFrame]) -> bytes:
    """abas: {'Nome da aba': DataFrame}. Devolve os bytes do .xlsx."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for nome, df in abas.items():
            nome = nome[:31]
            df.to_excel(xw, sheet_name=nome, index=False)
            _formatar(xw.sheets[nome], df)
    return buf.getvalue()

# -*- coding: utf-8 -*-
"""Reconhece e le os relatorios em Excel de verdade (.xlsx).

Em out/2026 o Protheus passou a exportar o rmatr029 em .xlsx (antes era XML
do Excel). Este modulo reconhece os relatorios em .xlsx pelo CONTEUDO
(cabecalho), nao pelo nome do arquivo. A distribuicao continua sendo a
planilha grande com a aba "base".
"""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import datetime

import pandas as pd


def norm(s) -> str:
    """'Comprador ' -> 'COMPRADOR'; 'DATA DISTRIBUIÇÃO' -> 'DATA DISTRIBUICAO'."""
    s = "".join(c for c in unicodedata.normalize("NFD", str(s or ""))
                if unicodedata.category(c) != "Mn")
    return " ".join(s.upper().split())


def _wb(b: bytes):
    import openpyxl

    return openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)


def _achar_tabela(wb, criterio) -> tuple | None:
    """Procura, nas primeiras linhas de cada aba, um cabecalho que satisfaca
    criterio(set_de_colunas_normalizadas). Devolve (aba, indice_linha, cabecalho)."""
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(max_row=8, values_only=True)):
            cab = [norm(c) for c in row]
            if criterio(set(cab)):
                return ws, i, list(row)
    return None


def _e_dist_base(cols: set) -> bool:
    return "RESPONSAVEL" in cols and any("DISTRIBUI" in c for c in cols)


def _e_sc(cols: set) -> bool:
    return "NUM.SC" in cols and "QTD.SOLICITADA" in cols


def _e_pc(cols: set) -> bool:
    return "PEDIDO COMPRA" in cols and "QUANT ENTREG" in cols


def tipo_xlsx(b: bytes) -> str | None:
    """'sc' | 'pc' | 'dist' | None, pelo cabecalho de alguma aba."""
    try:
        wb = _wb(b)
    except Exception:  # noqa: BLE001
        return None
    try:
        base = next((s for s in wb.sheetnames if s.strip().lower() == "base"), None)
        if base is not None:
            primeira = next(wb[base].iter_rows(max_row=1, values_only=True), ())
            if _e_dist_base({norm(c) for c in primeira}):
                return "dist"
        for tipo, crit in (("sc", _e_sc), ("pc", _e_pc)):
            if _achar_tabela(wb, crit):
                return tipo
        return None
    finally:
        wb.close()


def dt_ref_parametros(b: bytes) -> datetime | None:
    """'Dt.Ref: 07/10/2026' (ou 'Emissao: ...') da aba de parametros do Protheus."""
    try:
        wb = _wb(b)
    except Exception:  # noqa: BLE001
        return None
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_row=20, values_only=True):
                for c in row:
                    m = re.search(r"(?:DT\.REF|EMISSAO)[:\s]*([0-3]?\d/[0-1]?\d/\d{4})",
                                  norm(c))
                    if m:
                        return datetime.strptime(m.group(1), "%d/%m/%Y")
        return None
    finally:
        wb.close()


def ler_tabela(b: bytes, tipo: str) -> pd.DataFrame:
    """Le a tabela de dados (sc | pc) de um .xlsx, com o cabecalho original
    como nomes de coluna (sem espacos nas pontas)."""
    crit = {"sc": _e_sc, "pc": _e_pc}[tipo]
    wb = _wb(b)
    try:
        achado = _achar_tabela(wb, crit)
        if not achado:
            raise ValueError(f"Tabela de '{tipo}' nao encontrada no arquivo")
        ws, i_cab, cab = achado
        nomes, vistos = [], {}
        for c in cab:
            n = str(c).strip() if c is not None else ""
            if n in vistos:  # cabecalho repetido: 'X', 'X.1', ...
                vistos[n] += 1
                n = f"{n}.{vistos[n]}"
            else:
                vistos[n] = 0
            nomes.append(n)
        linhas = []
        for row in ws.iter_rows(min_row=i_cab + 2, values_only=True):
            if row is None or all(v is None or str(v).strip() == "" for v in row):
                continue
            row = list(row)[:len(nomes)] + [None] * max(0, len(nomes) - len(row))
            linhas.append(row)
        return pd.DataFrame(linhas, columns=nomes)
    finally:
        wb.close()
